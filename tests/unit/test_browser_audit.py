"""Offline regressions for browser lifecycle, truthful outcomes, and CV parsing."""

import asyncio
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from playwright.async_api import TimeoutError as PlaywrightTimeoutError

from agents.credentials_agent import CredentialsAgent
from agents.cv_parser_agent import CVParserAgent
from agents.form_filler_agent import FormFillerAgent
from agents.linkedin_agent import LinkedInAgent
from constants import TimingConstants
from utils.browser_setup import BrowserSetup
from utils.dom.dom_models import DOMElementNode
from utils.dom.dom_service import DomService


@pytest.fixture
def logs():
    return AsyncMock()


@pytest.fixture(autouse=True)
def short_delays(monkeypatch):
    for name in ('ACTION_DELAY', 'HUMAN_DELAY_MIN', 'FORM_FIELD_DELAY',
                 'FORM_SUBMIT_DELAY', 'PAGE_TRANSITION_DELAY', 'FILE_READ_DELAY',
                 'LLM_PROCESSING_DELAY', 'POLL_INTERVAL'):
        if name != 'POLL_INTERVAL':
            monkeypatch.setattr(TimingConstants, name, 0)


def make_linkedin(logs, page=None):
    controller = SimpleNamespace(settings={}, logs_manager=logs, tracker_agent=AsyncMock())
    page = page or AsyncMock()
    page.url = 'https://www.linkedin.com/jobs/'
    agent = LinkedInAgent(page, controller)
    agent._human_delay = AsyncMock()
    return agent


def make_form(logs, dom=None, **settings):
    settings['telemetry'] = {'enabled': False}
    agent = FormFillerAgent(dom or AsyncMock(), logs, settings)
    agent._human_delay = AsyncMock()
    return agent


def test_constructors_do_not_require_an_event_loop(logs):
    assert CVParserAgent({}, logs)
    assert DomService(MagicMock(), logs_manager=logs)


def test_linkedin_constructor_honors_settings_and_explicit_overrides(logs):
    controller = SimpleNamespace(logs_manager=logs, settings={'linkedin': {
        'default_timeout': 25000, 'min_delay': 1.2, 'max_delay': 2.4}})
    configured = LinkedInAgent(AsyncMock(), controller)
    assert (configured.default_timeout, configured.min_delay, configured.max_delay) == (25000, 1.2, 2.4)
    explicit = LinkedInAgent(AsyncMock(), controller, default_timeout=500, min_delay=0, max_delay=0)
    assert (explicit.default_timeout, explicit.min_delay, explicit.max_delay) == (500, 0, 0)


@pytest.mark.asyncio
@pytest.mark.parametrize('config', [{}, {'auto_apply_enabled': False}, {'auto_apply_enabled': 'true'}])
async def test_bulk_applications_require_explicit_boolean_opt_in(logs, config):
    agent = make_linkedin(logs)
    agent.settings['linkedin'] = config
    agent._handle_easy_apply = AsyncMock()
    assert await agent._apply_to_job({'is_easy_apply': True}) == 'skipped'
    assert await agent._apply_to_job({'is_easy_apply': False}) == 'skipped'
    agent._handle_easy_apply.assert_not_awaited()
    agent.page.query_selector.assert_not_awaited()
    agent.page.click.assert_not_awaited()


@pytest.mark.asyncio
async def test_bulk_application_runs_when_enabled(logs):
    agent = make_linkedin(logs)
    agent.settings['linkedin'] = {'auto_apply_enabled': True}
    agent._handle_easy_apply = AsyncMock(return_value='applied')
    assert await agent._apply_to_job({'job_title': 'Engineer', 'company': 'Example', 'is_easy_apply': True}) == 'applied'
    agent._handle_easy_apply.assert_awaited_once()


@pytest.mark.asyncio
@pytest.mark.parametrize('layout', ['.jobs-search-results-list [data-job-id]', '.jobs-job-board-list__item'])
@pytest.mark.parametrize('explicit_limit', [None, 100])
async def test_configured_job_limit_bounds_every_layout_and_persists_skips(logs, tmp_path, layout, explicit_limit):
    import csv
    agent = make_linkedin(logs)
    agent.settings['linkedin'] = {'job_search_limit': 2, 'auto_apply_enabled': False}
    agent.applied_jobs_csv = str(tmp_path / 'jobs.csv')
    cards = [AsyncMock() for _ in range(4)]
    for index, card in enumerate(cards):
        card.get_attribute.return_value = str(index)

    async def query(selector):
        return cards if selector == layout else []

    agent.page.query_selector_all.side_effect = query
    agent._extract_job_details = AsyncMock(return_value={'job_title': 'Engineer', 'company': 'Example'})
    assert await agent.process_job_listings(explicit_limit) == 2
    assert sum(card.click.await_count for card in cards) == 2
    with open(agent.applied_jobs_csv, newline='', encoding='utf-8') as stream:
        records = list(csv.DictReader(stream))
    assert len(records) == 2
    assert all(record['application_status'] == 'skipped' for record in records)


@pytest.mark.asyncio
async def test_failed_listing_is_persisted(logs, tmp_path):
    import csv
    agent = make_linkedin(logs)
    agent.settings['linkedin'] = {'job_search_limit': 1}
    agent.applied_jobs_csv = str(tmp_path / 'jobs.csv')
    card = AsyncMock()
    card.get_attribute.return_value = 'job-1'
    agent.page.query_selector_all.return_value = [card]
    agent.page.wait_for_selector.side_effect = PlaywrightTimeoutError('no details')
    assert await agent.process_job_listings() == 1
    with open(agent.applied_jobs_csv, newline='', encoding='utf-8') as stream:
        assert next(csv.DictReader(stream))['application_status'] == 'failed'


@pytest.mark.asyncio
@pytest.mark.parametrize('configured, expected', [(None, 50), (3, 3)])
async def test_search_passes_configured_limit_to_listing_processor(logs, configured, expected):
    agent = make_linkedin(logs)
    if configured is not None:
        agent.settings['linkedin'] = {'job_search_limit': configured}
    agent.check_captcha_or_logout = AsyncMock()
    agent._is_narrow_layout = AsyncMock(return_value=False)
    agent.process_job_listings = AsyncMock()
    await agent.search_jobs_and_apply('Engineer', 'Remote')
    agent.process_job_listings.assert_awaited_once_with(max_jobs=expected)


@pytest.mark.asyncio
@pytest.mark.parametrize('missing', ['job title', 'location'])
async def test_partial_standard_search_never_submits_or_processes_jobs(logs, missing):
    agent = make_linkedin(logs)
    agent.check_captcha_or_logout = AsyncMock()
    agent._is_narrow_layout = AsyncMock(return_value=False)
    agent.process_job_listings = AsyncMock()

    async def fill(selector, value):
        is_location = 'location' in selector.lower() or 'City, state' in selector
        if is_location == (missing == 'location'):
            raise PlaywrightTimeoutError('field unavailable')

    agent.page.fill.side_effect = fill
    with pytest.raises(RuntimeError, match=f'Could not set requested search criteria: {missing}'):
        await agent.search_jobs_and_apply('Engineer', 'Remote')
    agent.page.click.assert_not_awaited()
    agent.page.keyboard.press.assert_not_awaited()
    agent.process_job_listings.assert_not_awaited()


@pytest.mark.asyncio
@pytest.mark.parametrize('missing', ['job title', 'location'])
async def test_partial_responsive_search_only_expands_inputs(logs, missing):
    agent = make_linkedin(logs)
    icon = AsyncMock()
    agent.page.query_selector.return_value = icon

    async def fill(selector, value):
        is_location = 'location' in selector.lower() or 'City, state' in selector
        if is_location == (missing == 'location'):
            raise PlaywrightTimeoutError('field unavailable')

    agent.page.fill.side_effect = fill
    assert not await agent._handle_responsive_search('Engineer', 'Remote')
    icon.click.assert_awaited_once()  # Expand the controls, never submit.
    agent.page.query_selector.assert_awaited_once()
    agent.page.keyboard.press.assert_not_awaited()


@pytest.mark.asyncio
async def test_responsive_search_submits_when_both_fields_are_filled(logs):
    agent = make_linkedin(logs)
    icon, submit = AsyncMock(), AsyncMock()
    agent.page.query_selector.side_effect = [icon, submit]
    assert await agent._handle_responsive_search('Engineer', 'Remote')
    submit.click.assert_awaited_once()
    values = [call.args[1] for call in agent.page.fill.await_args_list]
    assert 'Engineer' in values and 'Remote' in values


@pytest.mark.asyncio
@pytest.mark.parametrize('title, location', [('', 'Remote'), ('Engineer', '  ')])
async def test_search_rejects_empty_criteria_before_page_actions(logs, title, location):
    agent = make_linkedin(logs)
    with pytest.raises(ValueError, match='nonempty job title and location'):
        await agent.search_jobs_and_apply(title, location)
    assert not await agent._handle_responsive_search(title, location)
    agent.page.query_selector.assert_not_awaited()
    agent.page.fill.assert_not_awaited()


@pytest.mark.asyncio
async def test_single_feed_extraction_obeys_configured_limit(logs):
    agent = make_linkedin(logs)
    agent.settings['linkedin'] = {'job_search_limit': 2}
    agent.page.query_selector_all.return_value = [AsyncMock() for _ in range(4)]
    agent._safe_get_text = AsyncMock(return_value='Example')
    assert len(await agent._handle_single_feed_layout(max_jobs=100)) == 2


@pytest.mark.asyncio
async def test_bundled_chromium_launch_has_no_chrome_channel(logs, tmp_path, monkeypatch):
    import utils.browser_setup as module
    page = AsyncMock()
    context = MagicMock(pages=[page], close=AsyncMock())
    launch = AsyncMock(return_value=context)
    driver = SimpleNamespace(stop=AsyncMock(), chromium=SimpleNamespace(launch_persistent_context=launch))
    monkeypatch.setattr(module, 'async_playwright', lambda: SimpleNamespace(start=AsyncMock(return_value=driver)))
    setup = BrowserSetup({'browser': {'type': 'chromium', 'data_dir': str(tmp_path)}, 'telemetry': {'enabled': False}}, logs)
    setup._configure_page = AsyncMock()
    setup.save_cookies = AsyncMock()
    resource, result = await setup.initialize()
    assert result is page
    assert launch.await_args.kwargs['channel'] is None
    await setup.cleanup(resource, result)
    driver.stop.assert_awaited_once()


@pytest.mark.asyncio
async def test_missing_dom_element_is_not_present():
    page = AsyncMock()
    page.wait_for_selector.side_effect = PlaywrightTimeoutError('missing')
    dom = DomService(page)
    assert await dom.wait_for_selector('#missing', timeout=1) is None
    assert not await dom.check_element_present('#missing', timeout=1)


@pytest.mark.asyncio
async def test_dom_tree_conversion_and_clickables_are_awaited():
    page = AsyncMock()
    page.evaluate.return_value = {'type': 'element', 'tag': 'button',
                                  'isClickable': True, 'isVisible': True}
    dom = DomService(page)
    tree = await dom.get_dom_tree()
    assert isinstance(tree, DOMElementNode)
    assert [element.tag for element in await dom.get_clickable_elements()] == ['button']


@pytest.mark.asyncio
async def test_repeated_dom_logging_bridge_does_not_reexpose(logs):
    page = AsyncMock()
    dom = DomService(page, logs_manager=logs)
    await dom._inject_logging_bridge()
    await dom._inject_logging_bridge()
    await DomService(page, logs_manager=logs)._inject_logging_bridge()
    page.expose_function.assert_awaited_once()


@pytest.mark.asyncio
async def test_failed_drag_releases_mouse_button():
    page = AsyncMock()
    source, target = AsyncMock(), AsyncMock()
    page.wait_for_selector.side_effect = [source, target]
    target.hover.side_effect = RuntimeError('detached')
    with pytest.raises(RuntimeError, match='detached'):
        await DomService(page).drag_and_drop('#source', '#target')
    page.mouse.up.assert_awaited_once()


@pytest.mark.asyncio
@pytest.mark.parametrize('method', ['verify_login_status', 'verify_captcha_success', 'verify_form_submission'])
async def test_missing_success_indicator_never_verifies(logs, method):
    dom = AsyncMock()
    dom.wait_for_selector.return_value = None
    agent = CredentialsAgent({}, dom, logs)
    assert not await getattr(agent, method)()
    selector = dom.wait_for_selector.await_args.args[0]
    assert isinstance(selector, str)
    assert 'div:not(' not in selector


@pytest.mark.asyncio
async def test_captcha_requests_use_https_timeouts_and_do_not_log_key(logs, monkeypatch):
    import agents.credentials_agent as module
    monkeypatch.setenv('TWO_CAPTCHA_API_KEY', 'private-test-key')
    post = MagicMock(side_effect=RuntimeError('https://host/?key=private-test-key'))
    monkeypatch.setattr(module.requests, 'post', post)
    agent = CredentialsAgent({}, AsyncMock(), logs)
    assert await agent._upload_to_2captcha(b'image') is None
    assert post.call_args.args[0].startswith('https://')
    assert post.call_args.kwargs['timeout'] == 30
    assert 'private-test-key' not in str(logs.mock_calls)


@pytest.mark.asyncio
async def test_cv_llm_defaults_preserve_extracted_text_and_filename(logs, tmp_path):
    path = tmp_path / 'resume.TXT'
    path.write_text('Real candidate experience', encoding='utf-8')
    agent = CVParserAgent({'use_llm': True}, logs)
    cv = await agent.parse_cv(path)
    assert cv.raw_text == 'Real candidate experience'
    assert cv.filename == 'resume.TXT'


@pytest.mark.asyncio
async def test_cv_cache_invalidated_when_file_changes(logs, tmp_path):
    path = tmp_path / 'resume.txt'
    path.write_text('Original', encoding='utf-8')
    agent = CVParserAgent({}, logs)
    _, first = await agent.prepare_cv(path)
    path.write_text('Updated experience', encoding='utf-8')
    assert agent.get_cached_cv(path) is None
    _, second = await agent.prepare_cv(path)
    assert first.raw_text == 'Original'
    assert second.raw_text == 'Updated experience'


@pytest.mark.asyncio
async def test_cv_empty_input_and_directories_rejected(logs, tmp_path):
    path = tmp_path / 'empty.txt'
    path.write_text('', encoding='utf-8')
    agent = CVParserAgent({}, logs)
    with pytest.raises(ValueError):
        await agent.parse_cv(path)
    assert not await agent.validate_for_upload(tmp_path)


@pytest.mark.asyncio
async def test_docx_text_is_parsed(logs, tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph('Candidate experience')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Python'
    path = tmp_path / 'resume.docx'
    document.save(path)
    cv = await CVParserAgent({}, logs).parse_cv(path)
    assert 'Candidate experience' in cv.raw_text
    assert 'Python' in cv.raw_text


@pytest.mark.asyncio
async def test_fill_form_returns_failure_for_missing_required_data(logs):
    agent = make_form(logs)
    assert not await agent.fill_form({}, {'name': {'selector': '#name', 'required': True}})
    agent.dom_service.wait_for_selector.assert_not_awaited()


@pytest.mark.asyncio
async def test_fill_form_accepts_dict_and_reports_success(logs):
    agent = make_form(logs)
    assert await agent.fill_form({'name': 'Candidate'}, {'name': {'selector': '#name'}})
    agent.dom_service.wait_for_selector.return_value.type.assert_awaited_once_with('Candidate')


@pytest.mark.asyncio
@pytest.mark.parametrize('confirmation, expected', [(None, False), (object(), True)])
async def test_submit_requires_positive_confirmation(logs, confirmation, expected):
    dom = AsyncMock()
    dom.query_selector.return_value = None
    button = AsyncMock()
    dom.wait_for_selector.side_effect = [button, confirmation]
    agent = make_form(logs, dom)
    assert await agent.submit_form('#submit') is expected
    button.click.assert_awaited_once()


@pytest.mark.asyncio
async def test_stale_confirmation_cannot_mark_another_application_as_sent(logs):
    dom = AsyncMock()
    dom.query_selector.return_value.is_visible.return_value = True
    agent = make_form(logs, dom)
    assert not await agent.submit_form('#submit')
    dom.wait_for_selector.assert_not_awaited()


@pytest.mark.asyncio
async def test_easy_apply_fills_final_step_before_submitting(logs):
    agent = make_form(logs)
    calls = []

    async def fill(_):
        calls.append('fill')
        return True

    async def submit(_):
        calls.append('submit')
        return False

    agent._handle_cv_upload = AsyncMock(return_value=True)
    agent._fill_current_step_fields = fill
    agent.submit_form = submit
    assert await agent.fill_easy_apply() == 'failed'
    assert calls == ['fill', 'submit']


@pytest.mark.asyncio
async def test_easy_apply_steps_are_bounded(logs):
    dom = AsyncMock()
    next_button = AsyncMock()

    async def query(selector):
        return next_button if 'Continue to next step' in selector else None

    dom.query_selector.side_effect = query
    agent = make_form(logs, dom, max_application_steps=3)
    agent._handle_cv_upload = AsyncMock(return_value=True)
    agent._fill_current_step_fields = AsyncMock(return_value=True)
    assert await agent.fill_easy_apply() == 'failed'
    assert next_button.click.await_count == 3


@pytest.mark.asyncio
async def test_required_checkbox_is_not_checked_without_explicit_answer(logs):
    dom = AsyncMock()
    dom.query_selector.return_value = None
    checkbox = AsyncMock()
    checkbox.is_visible.return_value = True
    checkbox.is_checked.return_value = False
    checkbox.get_attribute.return_value = 'qualified'
    dom.query_selector_all.return_value = [checkbox]
    agent = make_form(logs, dom)
    assert not await agent._fill_current_step_fields({})
    checkbox.click.assert_not_awaited()
    checkbox.check.assert_not_awaited()


@pytest.mark.asyncio
async def test_unique_cover_letter_files(logs):
    from pathlib import Path
    agent = make_form(logs)
    first = Path(await agent._write_cover_letter_to_file('first'))
    second = Path(await agent._write_cover_letter_to_file('second'))
    try:
        assert first != second
        assert first.read_text(encoding='utf-8') == 'first'
    finally:
        first.unlink()
        second.unlink()


@pytest.mark.asyncio
async def test_timing_contexts_support_async_with(logs):
    agent = make_linkedin(logs)
    async with agent._timed_operation('test'):
        pass
    async with agent._monitor_operation('test'):
        pass
    agent.controller.tracker_agent.log_activity.assert_awaited_once()


@pytest.mark.asyncio
async def test_linkedin_url_rejects_lookalike_host(logs):
    agent = make_linkedin(logs)
    agent.page.url = 'https://evil.example/linkedin.com/jobs/'
    assert not await agent._verify_url_is_jobs()
    agent.page.url = 'https://www.linkedin.com/jobs/search/?keywords=python'
    assert await agent._verify_url_is_jobs()


@pytest.mark.asyncio
async def test_two_column_listing_is_processed_once_and_loop_stops(logs):
    agent = make_linkedin(logs)
    card = AsyncMock()
    card.get_attribute.return_value = 'job-1'
    agent.page.query_selector_all.return_value = [card]
    agent._extract_job_details = AsyncMock(return_value={'job_title': 'Engineer'})
    agent._apply_to_job = AsyncMock(return_value='skipped')
    agent._save_job_record = AsyncMock()
    assert await agent.process_job_listings(max_jobs=10) == 1
    agent._apply_to_job.assert_awaited_once()


@pytest.mark.asyncio
async def test_attached_browser_uses_existing_context_and_preserves_browser(logs, tmp_path, monkeypatch):
    import utils.browser_setup as module
    page = AsyncMock()
    context = MagicMock(new_page=AsyncMock(return_value=page))
    browser = MagicMock(contexts=[context], new_page=AsyncMock(), close=AsyncMock())
    driver = SimpleNamespace(stop=AsyncMock(), chromium=SimpleNamespace(connect_over_cdp=AsyncMock(return_value=browser)))
    starter = SimpleNamespace(start=AsyncMock(return_value=driver))
    monkeypatch.setattr(module, 'async_playwright', lambda: starter)
    setup = BrowserSetup({'browser': {'type': 'chrome', 'attach_existing': True, 'data_dir': str(tmp_path)}, 'telemetry': {'enabled': False}}, logs)
    setup._configure_page = AsyncMock()
    resource, result = await setup.initialize()
    assert result is page
    context.new_page.assert_awaited_once()
    browser.new_page.assert_not_awaited()
    await setup.cleanup(resource, page)
    page.close.assert_awaited_once()
    browser.close.assert_not_awaited()
    driver.stop.assert_awaited_once()
    await setup.cleanup(None, None)
    driver.stop.assert_awaited_once()


@pytest.mark.asyncio
async def test_cancelled_browser_start_stops_driver(logs, tmp_path, monkeypatch):
    import utils.browser_setup as module
    driver = SimpleNamespace(stop=AsyncMock())
    monkeypatch.setattr(module, 'async_playwright', lambda: SimpleNamespace(start=AsyncMock(return_value=driver)))
    setup = BrowserSetup({'browser': {'type': 'firefox', 'data_dir': str(tmp_path)}, 'telemetry': {'enabled': False}}, logs)
    setup._launch_firefox_persistent = AsyncMock(side_effect=asyncio.CancelledError)
    with pytest.raises(asyncio.CancelledError):
        await setup.initialize()
    driver.stop.assert_awaited_once()


@pytest.mark.asyncio
async def test_cancelled_cookie_save_still_closes_owned_browser(logs, tmp_path):
    setup = BrowserSetup({'browser': {'data_dir': str(tmp_path)}, 'telemetry': {'enabled': False}}, logs)
    resource, driver = AsyncMock(), AsyncMock()
    setup._browser_or_context = resource
    setup._playwright = driver
    setup.save_cookies = AsyncMock(side_effect=asyncio.CancelledError)
    with pytest.raises(asyncio.CancelledError):
        await setup.cleanup(resource, AsyncMock())
    resource.close.assert_awaited_once()
    driver.stop.assert_awaited_once()


@pytest.mark.parametrize('browser_name,channel', [('edge', 'msedge'), ('chrome', 'chrome')])
async def test_normalized_browser_settings_preserve_requested_channel(logs, tmp_path, monkeypatch, browser_name, channel):
    import utils.browser_setup as module
    page = AsyncMock()
    context = SimpleNamespace(pages=[page])
    launch = AsyncMock(return_value=context)
    driver = SimpleNamespace(chromium=SimpleNamespace(launch_persistent_context=launch))
    monkeypatch.setattr(module, 'async_playwright', lambda: SimpleNamespace(start=AsyncMock(return_value=driver)))
    setup = BrowserSetup({'browser': {'type': 'chromium', 'raw_type': browser_name, 'should_prompt': False,
                                     'data_dir': str(tmp_path)}, 'telemetry': {'enabled': False}}, logs)
    setup._configure_page = AsyncMock()
    setup._get_browser_path = MagicMock(return_value=None)
    await setup.initialize()
    assert setup.browser_type == browser_name
    assert launch.call_args.kwargs['channel'] == channel
